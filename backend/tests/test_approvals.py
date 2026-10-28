"""Registration approval, "Confirm each registration" (app/api/approvals,
docs/APPROVALS.md): requests from the form, the club's queue and its
decisions, paid holds in the waiting-list machinery, and the lifecycle
(cancellation, the event over, retention, the guest merge).

Concurrency (two approvers, an approval racing a registration) is run for
real, on committed connections, in tests/test_approvals_races.py."""
import io
import time
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from types import SimpleNamespace

import pytest
import pytest_asyncio
from sqlalchemy import func, insert, select, text, update

from app.api.approvals import service as approvals
from app.api.approvals.models import RegistrationRequests as R
from app.api.events.models import EventRegistrationsLink, Events
from app.api.payments import service as payments
from app.api.ticketing import sweep
from app.api.ticketing.models import TicketTypes, WaitlistEntries
from app.api.users.models import SigninProviders
from app.core import jobs
from tests.conftest import auth_header
from tests.test_money_integrity import FakeRazorpay, rzp_fake  # noqa: F401
from tests.test_ticketing import make_code, set_types

API = "/api/v1"
NOW = datetime.now(timezone.utc)


@pytest.fixture
def mail(monkeypatch):
    from app.core.email import email as email_mod

    sent = []
    monkeypatch.setattr(email_mod, "send_email", lambda **kw: sent.append(kw))
    return sent


@pytest.fixture
def pushed(monkeypatch):
    from app.core.notifications import service as push

    got = []

    async def send(session, user_id, title, body, data=None):
        got.append((user_id, title, data))
        return True

    monkeypatch.setattr(push, "send_notification_to_user", send)
    return got


async def _shop(session, make_club, make_event, mint_token, *, tag="appr", **event_values):
    owner, club = await make_club(name=f"{tag.title()} Club", email=f"{tag}-owner@example.com")
    values = {"event_datetime": NOW + timedelta(days=5), "requires_approval": True, **event_values}
    event = await make_event(club=club, created_by=owner, name=f"{tag.title()} Night", **values)
    await session.commit()
    return SimpleNamespace(owner=owner, club=club, event=event, h=auth_header(await mint_token(owner)))


@pytest_asyncio.fixture
async def shop(session, make_club, make_event, mint_token, mail):
    w = await _shop(session, make_club, make_event, mint_token)
    w.mail = mail
    return w


def ask(client, event, n, headers=None, **extra):
    body = {
        "full_name": f"Asker Number {n}", "email": f"asker{n}@example.com", "phone": "98765" + f"{n:05d}"[-5:],
        "accepts_approval": True,
    }
    body.update(extra)
    return client.post(f"{API}/events/registration/{event.id}/register", json=body, headers=headers or {})


async def _req(session, rid) -> R:
    return await session.scalar(select(R).where(R.id == uuid.UUID(str(rid))).execution_options(populate_existing=True))


async def _regs(session, event_id) -> int:
    return int(await session.scalar(
        select(func.count()).select_from(EventRegistrationsLink).where(
            EventRegistrationsLink.event_id == event_id, EventRegistrationsLink.is_deleted == False,  # noqa: E712
        )
    ) or 0)


def approve(client, w, **body):
    return client.post(f"{API}/events/registration/{w.event.id}/requests/approve", json=body, headers=w.h)


def decline(client, w, **body):
    return client.post(f"{API}/events/registration/{w.event.id}/requests/decline", json=body, headers=w.h)


def listing(client, w, **params):
    return client.get(f"{API}/events/registration/{w.event.id}/requests", params=params, headers=w.h)


# ─── Submitting ──────────────────────────────────────────────────────────


async def test_a_request_is_made_once_and_the_form_sent_again_is_the_same_one(client, session, shop):
    r = await ask(client, shop.event, 1)
    assert r.status_code == 200, r.text
    body = r.json()
    a = body["approval"]
    assert (a["status"], body["event_registration_id"], body["event_id"]) == ("pending", None, shop.event.id)
    assert a["token"] and a["status_url"] == f"https://myotherapp.in/requests/{a['token']}"
    assert body["auth_token"]["access_token"]
    # Nobody is registered: a request holds no place.
    assert await _regs(session, shop.event.id) == 0
    req = await _req(session, a["id"])
    assert (req.full_name, req.email, req.status, req.whatsapp_opt_in) == ("Asker Number 1", "asker1@example.com", "pending", False)
    assert req.ticket_type_id is not None and req.user_id is not None
    # "You're on the waiting list", with the status link, once the request is kept.
    got = [m for m in shop.mail if m["recipients"] == ["asker1@example.com"]]
    assert len(got) == 1 and "waiting list" in got[0]["subject"] and a["token"] in got[0]["html"]

    # The form again, signed out: the same request. Whoever typed the address
    # gets no token; the link goes to the address instead (once in ten minutes).
    again = (await ask(client, shop.event, 1, email=" ASKER1@example.com")).json()
    assert again["approval"]["id"] == a["id"] and again["approval"]["token"] is None and again["approval"]["status_url"] is None
    assert again.get("auth_token") is None
    assert len([m for m in shop.mail if m["recipients"] == ["asker1@example.com"]]) == 2
    await ask(client, shop.event, 1)
    assert len([m for m in shop.mail if m["recipients"] == ["asker1@example.com"]]) == 2
    # With the guest's own session it is theirs: the token comes back.
    mine = (await ask(client, shop.event, 1, headers=auth_header(body["auth_token"]["access_token"]))).json()
    assert mine["approval"]["id"] == a["id"] and mine["approval"]["token"] == a["token"]
    assert await session.scalar(select(func.count()).select_from(R).where(R.event_id == shop.event.id)) == 1

    # The event says so, and the guest's page and list show it.
    info = (await client.get(f"{API}/events/info/{shop.event.id}", headers=auth_header(body["auth_token"]["access_token"]))).json()
    assert info["requires_approval"] is True and info["ticketing"]["requires_approval"] is True
    assert info["ticketing"]["viewer_request"] == {"id": a["id"], "status": "pending", "token": a["token"], "status_url": a["status_url"]}
    page = await client.get(f"{API}/registration-requests/{a['token']}")
    assert page.status_code == 200 and page.headers["cache-control"] == "no-store"
    p = page.json()
    assert (p["status"], p["event"]["name"], p["event"]["club"]["name"], p["ticket_id"], p["pay"]) == ("pending", "Appr Night", "Appr Club", None, None)
    listed = (await client.get(f"{API}/registration-requests/mine", headers=auth_header(body["auth_token"]["access_token"]))).json()
    assert [i["id"] for i in listed["items"]] == [a["id"]] and listed["items"][0]["token"] == a["token"]
    assert (await client.get(f"{API}/registration-requests/not-a-real-token-at-all")).status_code == 404


async def test_an_old_client_is_told_to_update(client, session, shop):
    page = await client.get(f"{API}/discover/events/{shop.event.slug}")
    assert page.status_code == 200 and page.json()["event"]["requires_approval"] is True
    old = await ask(client, shop.event, 2, accepts_approval=False)
    assert old.status_code == 409 and old.json()["error_code"] == "APPROVAL_REQUIRED"
    legacy = await client.post(
        f"{API}/events/registration/{shop.event.id}/register",
        json={"full_name": "Old App User", "email": "oldapp@example.com"},
    )
    assert legacy.status_code == 409 and legacy.json()["error_code"] == "APPROVAL_REQUIRED"
    assert await session.scalar(select(func.count()).select_from(R).where(R.event_id == shop.event.id)) == 0
    # Picks from the shelf are not taken with a request.
    picks = await ask(client, shop.event, 2, store=[{"item_id": str(uuid.uuid4()), "quantity": 1}])
    assert picks.status_code == 422 and picks.json()["error_code"] == "ADD_ONS_NOT_ALLOWED"
    # The setting off: the form registers directly again.
    await session.execute(update(Events).where(Events.id == shop.event.id).values(requires_approval=False))
    await session.commit()
    direct = await ask(client, shop.event, 2, accepts_approval=False)
    assert direct.status_code == 200 and direct.json()["event_registration_id"]


async def test_declined_is_final_withdrawn_may_ask_again(client, session, shop):
    first = (await ask(client, shop.event, 3)).json()["approval"]
    d = (await decline(client, shop, ids=[first["id"]], notify=False)).json()
    assert d == {"declined": [first["id"]], "skipped": []}
    again = await ask(client, shop.event, 3)
    assert again.status_code == 409 and again.json()["error_code"] == "REQUEST_DECLINED"
    # Declined can't be withdrawn: the request as it stands comes back with the 409.
    w = await client.post(f"{API}/registration-requests/{first['token']}/withdraw")
    assert w.status_code == 409 and w.json()["error_code"] == "REQUEST_NOT_WITHDRAWABLE"
    assert w.json()["errors"]["request"]["status"] == "declined"
    # The club reopens it: back in the queue, where it was.
    r = (await client.post(f"{API}/events/registration/{shop.event.id}/requests/reopen", json={"ids": [first["id"]]}, headers=shop.h)).json()
    assert r == {"reopened": [first["id"]], "skipped": []}
    req = await _req(session, first["id"])
    assert (req.status, req.decided_at, req.decline_note) == ("pending", None, None)

    # Withdrawn (twice is the same), then asked for again: a new request.
    gone = await client.post(f"{API}/registration-requests/{first['token']}/withdraw")
    assert gone.status_code == 200 and gone.json()["status"] == "withdrawn"
    assert (await client.post(f"{API}/registration-requests/{first['token']}/withdraw")).json()["status"] == "withdrawn"
    new = (await ask(client, shop.event, 3)).json()["approval"]
    assert new["id"] != first["id"] and new["status"] == "pending" and new["token"]
    # A withdrawn one is not reopened over the open one.
    r = (await client.post(f"{API}/events/registration/{shop.event.id}/requests/reopen", json={"ids": [first["id"]]}, headers=shop.h)).json()
    assert r["skipped"] == [{"id": first["id"], "reason": "not_reopenable"}]


async def test_the_list_pages_by_key_counts_and_searches(client, session, shop):
    made = [(await ask(client, shop.event, 10 + i)).json()["approval"]["id"] for i in range(7)]
    await decline(client, shop, ids=[made[5]], note="Full up")
    page1 = (await listing(client, shop, limit=3)).json()
    assert [i["id"] for i in page1["items"]] == made[:3] and page1["next"]
    assert page1["counts"] == {"pending": 6, "approved": 0, "declined": 1, "withdrawn": 0, "lapsed": 0, "closed": 0, "total": 7}
    assert page1["places_left"] is None
    item = page1["items"][0]
    assert (item["full_name"], item["email"], item["status"], item["ticket_type"]["name"]) == ("Asker Number 10", "asker10@example.com", "pending", "General admission")
    page2 = (await listing(client, shop, limit=3, after=page1["next"])).json()
    assert [i["id"] for i in page2["items"]] == made[3:5] + [made[6]] and page2["next"] is None
    two = (await listing(client, shop, limit=2, after=page1["next"])).json()
    assert [i["id"] for i in two["items"]] == made[3:5] and two["next"]
    assert [i["id"] for i in (await listing(client, shop, limit=2, after=two["next"])).json()["items"]] == [made[6]]
    found = (await listing(client, shop, q="asker12@")).json()
    assert [i["id"] for i in found["items"]] == [made[2]]
    declined = (await listing(client, shop, status="declined")).json()["items"]
    assert [(i["id"], i["decline_note"], i["decided_by"]["name"]) for i in declined] == [(made[5], "Full up", "Test User")]
    bad = await listing(client, shop, after="not-a-cursor")
    assert bad.status_code == 422
    # Door staff read, never decide (club_access.policy).
    from app.api.club_access import policy

    assert policy.allows("door", "GET", f"/api/v1/events/registration/{shop.event.id}/requests")
    assert not policy.allows("door", "POST", f"/api/v1/events/registration/{shop.event.id}/requests/approve")
    assert policy.allows("editor", "POST", f"/api/v1/events/registration/{shop.event.id}/requests/approve")
    assert policy.can("door")["decide_requests"] is False and policy.can("editor")["decide_requests"] is True


async def test_the_club_switches_it_on_and_off_from_the_event_form(client, session, make_club, mint_token):
    from app.api.events.models import EventCategories

    owner, club = await make_club(name="Form Club", email="formclub@example.com")
    category = EventCategories(name="Form things")
    session.add(category)
    await session.commit()
    h = auth_header(await mint_token(owner))
    form = {
        "name": "Form Night", "event_datetime": (NOW + timedelta(days=9)).isoformat(), "duration": "2",
        "has_fee": "false", "status": "published", "category_id": str(category.id),
    }
    made = await client.post(f"{API}/events/create", data={**form, "requires_approval": "true"}, headers=h)
    assert made.status_code in (200, 201), made.text
    assert made.json()["requires_approval"] is True
    event_id = made.json()["id"]
    # Left out of an edit: as it was.
    kept = await client.put(f"{API}/events/update/{event_id}", data=form, headers=h)
    assert kept.status_code == 200 and kept.json()["requires_approval"] is True
    off = await client.put(f"{API}/events/update/{event_id}", data={**form, "requires_approval": "false"}, headers=h)
    assert off.json()["requires_approval"] is False
    assert (await client.get(f"{API}/events/info/{event_id}")).json()["requires_approval"] is False
    plain = await client.post(f"{API}/events/create", data={**form, "name": "Plain Night"}, headers=h)
    assert plain.json()["requires_approval"] is False


# ─── Confirming ──────────────────────────────────────────────────────────


async def _bulk_requests(session, event, n: int, *, tag: str, type_id: int) -> list:
    """n pending requests, written straight in (the form is tested above)."""
    now = datetime.now(timezone.utc)
    rows = [
        dict(
            id=uuid.uuid4(), event_id=event.id, user_id=None, full_name=f"Bulk {tag} {i}", email=f"bulk-{tag}-{i}@example.com",
            ticket_type_id=type_id, status="pending", token=approvals.new_token(), submitted_at=now + timedelta(microseconds=i),
            whatsapp_opt_in=False, club_updates_whatsapp=False, created_at=now, updated_at=now,
        )
        for i in range(n)
    ]
    await session.execute(insert(R), rows)
    await session.commit()
    return [r["id"] for r in rows]


async def _default_type(session, event) -> int:
    from app.api.ticketing import types as types_mod

    types = await types_mod.ensure(session, event)
    await session.commit()
    return types[0].id


async def test_five_hundred_confirmed_in_one_call(client, session, shop, monkeypatch, pushed):
    from app.api.deliveries import service as deliveries
    from app.api.deliveries.models import Deliveries

    # The tickets' sends are the outbox's business (tested there): left queued here.
    async def nothing(*_a, **_k):
        return 0

    monkeypatch.setattr(deliveries, "_run_due", nothing)
    ids = await _bulk_requests(session, shop.event, 500, tag="five", type_id=await _default_type(session, shop.event))
    started = time.monotonic()
    r = await approve(client, shop, ids=[str(i) for i in ids])
    took = time.monotonic() - started
    assert r.status_code == 200, r.text
    out = r.json()
    assert len(out["approved"]) == 500 and out["skipped"] == [] and out["places_left"] is None
    assert [a["id"] for a in out["approved"]] == [str(i) for i in ids]
    assert all(a["registration_id"] and a["pay_by"] is None for a in out["approved"])
    assert await _regs(session, shop.event.id) == 500
    statuses = dict((await session.execute(select(R.status, func.count()).where(R.event_id == shop.event.id).group_by(R.status))).all())
    assert statuses == {"approved": 500}
    # Each one's ticket, queued in the outbox with the decision.
    assert await session.scalar(select(func.count()).select_from(Deliveries).where(Deliveries.event_id == shop.event.id)) == 500
    going = await session.scalar(select(Events.going_count).where(Events.id == shop.event.id).execution_options(populate_existing=True))
    assert going == 500
    print(f"500 confirmations in {took:.1f}s")
    # Confirming the same again: nothing twice.
    again = (await approve(client, shop, ids=[str(i) for i in ids[:5]])).json()
    assert again["approved"] == [] and {s["reason"] for s in again["skipped"]} == {"not_pending"}
    too_many = await approve(client, shop, ids=[str(uuid.uuid4()) for _ in range(501)])
    assert too_many.status_code == 422


async def test_confirming_stops_at_the_event_and_the_types_places(client, session, make_club, make_event, mint_token, mail, pushed):
    w = await _shop(session, make_club, make_event, mint_token, tag="capped", max_participants=3)
    made = [(await ask(client, w.event, 40 + i)).json()["approval"]["id"] for i in range(5)]
    out = (await approve(client, w, all_pending=True, limit=500)).json()
    assert [a["id"] for a in out["approved"]] == made[:3]
    assert out["skipped"] == [{"id": made[3], "reason": "event_full"}, {"id": made[4], "reason": "event_full"}]
    assert out["places_left"] == 0 and await _regs(session, w.event.id) == 3
    # The ticket email, queued in the outbox for the three confirmed, nobody else.
    from app.api.deliveries.models import Deliveries

    queued = (await session.scalars(
        select(EventRegistrationsLink.email).join(Deliveries, Deliveries.registration_id == EventRegistrationsLink.id)
        .where(Deliveries.event_id == w.event.id)
    )).all()
    assert sorted(queued) == [f"asker{40 + i}@example.com" for i in range(3)]
    assert (await listing(client, w)).json()["counts"]["pending"] == 2

    # A type's own cap: its rest skipped, the other type's still confirmed.
    t = await _shop(session, make_club, make_event, mint_token, tag="typed")
    types = await set_types(client, SimpleNamespace(h=t.h), t.event, [{"name": "Seat", "capacity": 1}, {"name": "Lawn"}])
    seat, lawn = types["Seat"]["id"], types["Lawn"]["id"]
    a = (await ask(client, t.event, 50, ticket_type_id=seat)).json()["approval"]["id"]
    b = (await ask(client, t.event, 51, ticket_type_id=seat)).json()["approval"]["id"]
    c = (await ask(client, t.event, 52, ticket_type_id=lawn)).json()["approval"]["id"]
    out = (await approve(client, t, ids=[a, b, c])).json()
    assert [x["id"] for x in out["approved"]] == [a, c] and out["skipped"] == [{"id": b, "reason": "type_sold_out"}]
    assert await session.scalar(select(TicketTypes.taken_count).where(TicketTypes.id == seat).execution_options(populate_existing=True)) == 1
    # Pushes went to app accounts only: these were all guests.
    assert pushed == []


async def test_an_app_user_is_confirmed_and_told_in_the_app(client, session, shop, make_user, mint_token, pushed):
    user = await make_user(email="appasker@example.com", username="appasker")
    await session.commit()
    h = auth_header(await mint_token(user))
    a = (await ask(client, shop.event, 60, email="appasker@example.com", headers=h)).json()
    assert a["approval"]["token"] and a.get("auth_token") is None
    req = await _req(session, a["approval"]["id"])
    assert req.user_id == user.id
    out = (await approve(client, shop, ids=[a["approval"]["id"]])).json()
    reg_id = out["approved"][0]["registration_id"]
    reg = await session.get(EventRegistrationsLink, uuid.UUID(reg_id))
    assert reg.user_id == user.id and reg.full_name == "Asker Number 60"
    assert [(u, d["type"], d["token"], d["event_id"]) for u, _t, d in pushed] == [(user.id, "request_approved", req.token, str(shop.event.id))]
    view = (await client.get(f"{API}/registration-requests/{req.token}")).json()
    assert view["status"] == "approved" and view["ticket_id"] == reg.ticket_id
    info = (await client.get(f"{API}/events/info/{shop.event.id}", headers=h)).json()
    assert info["is_registered"] is True and info["ticketing"]["viewer_request"] is None
    # Registered: the form answers as it always did.
    dup = await ask(client, shop.event, 60, email="appasker@example.com", headers=h)
    assert dup.status_code == 400 and dup.json()["error_code"] == "ALREADY_REGISTERED"


async def test_decline_tells_them_only_when_asked_and_reopen_puts_it_back(client, session, shop, make_user, mint_token, pushed):
    user = await make_user(email="declined-app@example.com", username="declinedapp")
    await session.commit()
    told = (await ask(client, shop.event, 70, email="declined-app@example.com", headers=auth_header(await mint_token(user)))).json()["approval"]
    quiet = (await ask(client, shop.event, 71)).json()["approval"]
    shop.mail.clear()
    assert (await decline(client, shop, ids=[told["id"]], note="We're full up, {sorry}")).json()["declined"] == [told["id"]]
    assert (await decline(client, shop, ids=[quiet["id"]], notify=False)).json()["declined"] == [quiet["id"]]
    assert [m["recipients"] for m in shop.mail] == [["declined-app@example.com"]]
    assert "couldn't offer you a place" in shop.mail[0]["html"] and "(sorry)" in shop.mail[0]["html"]
    assert [(u, d["type"]) for u, _t, d in pushed] == [(user.id, "request_declined")]
    from app.api.notifications.models import Notifications

    inbox = await session.scalar(select(Notifications).where(Notifications.user_id == user.id, Notifications.type == "request_declined"))
    assert inbox is not None and inbox.event_id == shop.event.id
    # Declining again, or a decided one: skipped.
    again = (await decline(client, shop, ids=[told["id"]])).json()
    assert again == {"declined": [], "skipped": [{"id": told["id"], "reason": "not_pending"}]}
    view = (await client.get(f"{API}/registration-requests/{told['token']}")).json()
    assert view["status"] == "declined" and view["decline_note"] == "We're full up, {sorry}"
    r = (await client.post(f"{API}/events/registration/{shop.event.id}/requests/reopen", json={"ids": [told["id"], quiet["id"]]}, headers=shop.h)).json()
    assert r["reopened"] == [told["id"], quiet["id"]]
    assert (await listing(client, shop)).json()["counts"]["pending"] == 2
    # Every waiting one at once, oldest first, without paging through their ids.
    third = (await ask(client, shop.event, 72)).json()["approval"]["id"]
    shop.mail.clear()
    some = (await decline(client, shop, all_pending=True, limit=2, notify=False)).json()
    assert some == {"declined": [told["id"], quiet["id"]], "skipped": []}
    rest = (await decline(client, shop, all_pending=True)).json()
    assert rest["declined"] == [third] and [m["recipients"] for m in shop.mail] == [["asker72@example.com"]]
    assert (await decline(client, shop)).status_code == 422


async def test_emails_go_only_once_the_decision_is_kept(client, session, shop, monkeypatch):
    ids = [(await ask(client, shop.event, 80 + i)).json()["approval"]["id"] for i in range(2)]
    event_id = shop.event.id
    shop.mail.clear()
    before = await session.scalar(select(func.count()).select_from(jobs.Jobs).where(jobs.Jobs.name == approvals.NOTIFY_JOB))
    # A decline whose transaction is rolled back: nothing is written, nothing is sent.
    real_commit = session.commit

    async def roll_back():
        await session.rollback()

    monkeypatch.setattr(session, "commit", roll_back)
    await approvals.decline(session, await session.get(Events, event_id), ids)
    monkeypatch.setattr(session, "commit", real_commit)
    await jobs.settle()
    assert shop.mail == []
    assert await session.scalar(select(func.count()).select_from(jobs.Jobs).where(jobs.Jobs.name == approvals.NOTIFY_JOB)) == before
    assert {(await _req(session, i)).status for i in ids} == {"pending"}
    # Kept: the job is written with it, and the mail goes once it has committed.
    await approvals.decline(session, await session.get(Events, event_id), ids[:1])
    assert shop.mail == []
    await jobs.settle()
    assert [m["recipients"] for m in shop.mail] == [["asker80@example.com"]]


# ─── Paid tickets: a hold, a claim, a payment ─────────────────────────────


async def test_a_paid_confirmation_holds_the_place_until_paid_or_it_lapses(client, session, make_club, make_event, mint_token, mail, rzp_fake, pushed):
    w = await _shop(session, make_club, make_event, mint_token, tag="paidappr", has_fee=True, reg_fee=Decimal("199.00"), max_participants=2)
    ids = [(await ask(client, w.event, 90 + i, additional_details={"note": f"hi {i}"})).json()["approval"]["id"] for i in range(3)]
    mail.clear()
    out = (await approve(client, w, ids=ids[:2])).json()
    assert [a["id"] for a in out["approved"]] == ids[:2] and all(a["pay_by"] and a["registration_id"] is None for a in out["approved"])
    assert out["places_left"] == 0
    # Nobody registered yet; both holds are offers of the waiting-list machinery.
    assert await _regs(session, w.event.id) == 0
    holds = (await session.scalars(select(WaitlistEntries).where(WaitlistEntries.event_id == w.event.id).execution_options(populate_existing=True))).all()
    assert len(holds) == 2 and all(h.status == "offered" and h.via_request for h in holds)
    # ...but never the club's waiting list.
    wl = (await client.get(f"{API}/events/{w.event.id}/ticketing/waitlist", headers=w.h)).json()
    assert wl["entries"] == [] and wl["counts"]["offered"] == 0
    # The third: no place left while the two are held.
    third = (await approve(client, w, ids=ids[2:])).json()
    assert third["skipped"] == [{"id": ids[2], "reason": "event_full"}]
    # "Pay ₹199 by …", with the claim link.
    pay_mail = [m for m in mail if m["recipients"] == ["asker90@example.com"]]
    assert len(pay_mail) == 1 and "confirmed your place" in pay_mail[0]["subject"] and "₹199" in pay_mail[0]["html"]

    first = await _req(session, ids[0])
    view = (await client.get(f"{API}/registration-requests/{first.token}")).json()
    assert view["status"] == "approved" and view["pay"]["amount"] == 199.0 and view["ticket_id"] is None
    claim_token = view["pay"]["claim_token"]
    offer = (await client.get(f"{API}/waitlist/offers/{claim_token}")).json()
    assert (offer["reason"], offer["amount"], offer["event"]["club"]["name"], offer["needs_payment"]) == ("approved", 199.0, "Paidappr Club", True)
    # The claim registers what the request said, whatever the page sends.
    claimed = await client.post(f"{API}/waitlist/offers/{claim_token}/claim", json={"full_name": "Someone Else", "additional_details": {"x": "y"}})
    assert claimed.status_code == 200, claimed.text
    c = claimed.json()
    assert (c["full_name"], c["pay_amount"], c["is_paid"], c["additional_details"]) == ("Asker Number 90", 199.0, False, {"note": "hi 0"})
    assert c["auth_token"]
    first = await _req(session, ids[0])
    assert str(first.registration_id) == c["event_registration_id"] and first.status == "approved"
    order = (await client.post(
        f"{API}/payments/orders",
        json={"source": "event_registration", "payload": {"event_id": w.event.id, "event_registration_id": c["event_registration_id"]}},
        headers=auth_header(c["auth_token"]["access_token"]),
    )).json()
    p = rzp_fake.pay(order["razorpay_order_id"])
    await payments.verify_razorpay_payment(session, order["razorpay_order_id"], p["id"], payment_details=p, expand_payment_details=False)
    reg = await session.get(EventRegistrationsLink, uuid.UUID(c["event_registration_id"]))
    await session.refresh(reg)
    assert reg.is_paid is True and p["id"] not in rzp_fake.refunded()
    view = (await client.get(f"{API}/registration-requests/{first.token}")).json()
    assert view["status"] == "approved" and view["ticket_id"] == reg.ticket_id and view["pay"] is None

    # The second never pays: its hold runs out, the request lapses, the place goes back.
    second = await _req(session, ids[1])
    await session.execute(update(WaitlistEntries).where(WaitlistEntries.id == second.waitlist_entry_id).values(offer_expires_at=NOW - timedelta(minutes=1)))
    await session.commit()
    mail.clear()
    await sweep.run_sweep(session)
    await jobs.settle()
    second = await _req(session, ids[1])
    assert second.status == "lapsed"
    assert [m["recipients"] for m in mail] == [["asker91@example.com"]] and "released" in mail[0]["subject"]
    assert (await client.get(f"{API}/registration-requests/{second.token}")).json()["pay"] is None
    # The third gets it now.
    now = (await approve(client, w, ids=ids[2:])).json()
    assert [a["id"] for a in now["approved"]] == ids[2:] and now["places_left"] == 0
    # A lapsed one can be put back in the queue.
    r = (await client.post(f"{API}/events/registration/{w.event.id}/requests/reopen", json={"ids": [ids[1]]}, headers=w.h)).json()
    assert r["reopened"] == [ids[1]]


async def test_a_confirmed_paid_place_given_back_goes_on(client, session, make_club, make_event, mint_token, mail):
    w = await _shop(session, make_club, make_event, mint_token, tag="giveback", has_fee=True, reg_fee=Decimal("99.00"), max_participants=1)
    rid = (await ask(client, w.event, 95)).json()["approval"]["id"]
    await approve(client, w, ids=[rid])
    req = await _req(session, rid)
    gone = await client.post(f"{API}/registration-requests/{req.token}/withdraw")
    assert gone.status_code == 200 and gone.json()["status"] == "withdrawn"
    entry = await session.scalar(select(WaitlistEntries).where(WaitlistEntries.id == req.waitlist_entry_id).execution_options(populate_existing=True))
    assert entry.status == "left" and entry.offer_token is None
    assert (await listing(client, w)).json()["places_left"] == 1
    # "No thanks" on the claim page is the same.
    rid2 = (await ask(client, w.event, 96)).json()["approval"]["id"]
    await approve(client, w, ids=[rid2])
    req2 = await _req(session, rid2)
    token = (await session.get(WaitlistEntries, req2.waitlist_entry_id)).offer_token
    assert (await client.post(f"{API}/waitlist/offers/{token}/decline")).status_code == 200
    assert (await _req(session, rid2)).status == "withdrawn"


async def test_a_code_no_longer_good_is_skipped_for_the_club_to_decide(client, session, make_club, make_event, mint_token, mail):
    """Never confirmed at full price behind the person's back: a code that no
    longer applies skips the row; a code's last use is never promised twice."""
    w = await _shop(session, make_club, make_event, mint_token, tag="coded", has_fee=True, reg_fee=Decimal("200.00"))
    code = await make_code(client, SimpleNamespace(h=w.h), code="HALF", kind="percent", value=50, event_id=w.event.id, max_uses=1)
    a = (await ask(client, w.event, 150, discount_code=" half ")).json()["approval"]["id"]
    b = (await ask(client, w.event, 151, discount_code="HALF")).json()["approval"]["id"]
    assert (await _req(session, a)).discount_code == "HALF"
    off = await client.patch(f"{API}/clubs/discount-codes/{code['id']}", json={"is_active": False}, headers=w.h)
    assert off.status_code == 200, off.text
    out = (await approve(client, w, ids=[a])).json()
    assert out["skipped"] == [{"id": a, "reason": "code_invalid"}]
    await client.patch(f"{API}/clubs/discount-codes/{code['id']}", json={"is_active": True}, headers=w.h)
    out = (await approve(client, w, ids=[a, b])).json()
    # One use left: the first is held at the code's price, the second skipped.
    assert [x["id"] for x in out["approved"]] == [a] and out["skipped"] == [{"id": b, "reason": "code_invalid"}]
    assert (await _req(session, a)).amount == Decimal("100.00")
    token = (await client.get(f"{API}/registration-requests/{(await _req(session, a)).token}")).json()["pay"]["claim_token"]
    claimed = (await client.post(f"{API}/waitlist/offers/{token}/claim", json={})).json()
    assert claimed["pay_amount"] == 100.0 and claimed["discount"]["code"] == "HALF"


# ─── The event called off or over, retention, the guest merge ────────────


async def test_cancelling_closes_the_requests_and_tells_them(client, session, shop):
    ids = [(await ask(client, shop.event, 100 + i)).json()["approval"]["id"] for i in range(2)]
    shop.mail.clear()
    r = await client.post(f"{API}/events/{shop.event.id}/cancel", json={"reason": "The venue flooded"}, headers=shop.h)
    assert r.status_code == 200, r.text
    await jobs.settle()
    assert {(await _req(session, i)).status for i in ids} == {"closed"}
    cancelled = [m for m in shop.mail if m["subject"].startswith("Cancelled")]
    assert sorted(m["recipients"][0] for m in cancelled) == ["asker100@example.com", "asker101@example.com"]
    assert "The venue flooded" in cancelled[0]["html"]
    late = await ask(client, shop.event, 102)
    assert late.status_code == 400 and late.json()["error_code"] == "REGISTRATIONS_CLOSED"


async def test_an_event_over_has_its_pending_requests_closed(client, session, shop):
    rid = (await ask(client, shop.event, 110)).json()["approval"]["id"]
    await session.execute(update(Events).where(Events.id == shop.event.id).values(event_datetime=NOW - timedelta(days=1)))
    await session.commit()
    out = (await approve(client, shop, ids=[rid])).json()
    assert out["skipped"] == [{"id": rid, "reason": "event_closed"}]
    await sweep.run_sweep(session)
    assert (await _req(session, rid)).status == "closed"


async def test_the_retention_scrub_takes_the_requests_too(client, session, shop):
    from app.api.closeout import service as closeout

    rid = (await ask(client, shop.event, 120, additional_details={"college": "CET"})).json()["approval"]["id"]
    await decline(client, shop, ids=[rid], note="Sorry", notify=False)
    event = await session.get(Events, shop.event.id)
    job = await closeout.ensure_job(session, event)
    assert (await closeout.preview_counts(session, event))["requests"] == 1
    await closeout.scrub(session, event, job)
    req = await _req(session, rid)
    assert (req.full_name, req.phone, req.answers, req.decline_note) == ("Removed", None, None, None)
    assert req.email.startswith("removed-") and req.email.endswith("@removed.invalid") and "asker" not in req.email


async def test_the_guest_merge_moves_requests_to_the_account(client, session, shop, make_user):
    from app.api.users import service as users

    rid = (await ask(client, shop.event, 130, email="merger@example.com")).json()["approval"]["id"]
    guest_id = (await _req(session, rid)).user_id
    app_user = await make_user(email="Merger@example.com", username="merger")
    app_user.provider = SigninProviders.google
    await session.flush()
    summary = await users.merge_guest_into_user(session, app_user)
    await session.commit()
    assert summary["requests_transferred"] == 1
    req = await _req(session, rid)
    assert req.user_id == app_user.id != guest_id


# ─── Nothing else sees a request ─────────────────────────────────────────


async def test_the_door_certificates_and_exports_never_see_a_request(client, session, shop):
    from openpyxl import load_workbook

    from app.api.checkins import service as checkins

    for i in range(3):
        assert (await ask(client, shop.event, 140 + i)).status_code == 200
    assert (await client.get(f"{API}/events/registration/{shop.event.id}/list", headers=shop.h)).json()["total"] == 0
    exported = await client.get(f"{API}/events/registration/{shop.event.id}/export", headers=shop.h)
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert sheet.max_row <= 1
    event = await session.get(Events, shop.event.id)
    assert await checkins.search(session, event, "Asker") == []
    preview = (await client.post(f"{API}/events/{shop.event.id}/certificates/issue-preview", json={"audience": "everyone"}, headers=shop.h)).json()
    assert preview["eligible"] == 0
    summary = (await client.get(f"{API}/events/{shop.event.id}/summary", headers=shop.h)).json()
    assert summary["registered"] == 0 and summary["requests_pending"] == 3
    dash = (await client.get(f"{API}/clubs/dashboard", headers=shop.h)).json()
    assert [a["kind"] for a in dash["attention"]][:1] == ["REQUESTS_PENDING"]
    lib = await client.get(f"{API}/clubs/me/emails/overview", headers=shop.h)
    assert lib.status_code == 200, lib.text
    assert {e["key"] for e in lib.json()["automatic"]} >= {"request_received", "request_approved_pay", "request_declined", "request_lapsed"}


# ─── The queries use the index ───────────────────────────────────────────


async def test_the_queue_and_the_confirmation_read_the_new_index(session, make_club, make_event, mint_token):
    """EXPLAIN of the club's page, the counts and the confirmation's lock, on
    a table with other events' requests in it (recorded in the PR)."""
    w = await _shop(session, make_club, make_event, mint_token, tag="explain")
    others = [await make_event(club=w.club, created_by=w.owner, name=f"Other {i}", event_datetime=NOW + timedelta(days=9)) for i in range(20)]
    await session.commit()
    type_id = await _default_type(session, w.event)
    ids = await _bulk_requests(session, w.event, 300, tag="ex", type_id=type_id)
    for i, e in enumerate(others):
        await _bulk_requests(session, e, 200, tag=f"o{i}", type_id=await _default_type(session, e))
    await session.execute(text("ANALYZE registration_requests"))

    async def plan(stmt) -> str:
        compiled = stmt.compile(dialect=session.bind.dialect, compile_kwargs={"literal_binds": True})
        rows = (await session.execute(text(f"EXPLAIN {compiled}"))).scalars().all()
        return "\n".join(rows)

    page = await plan(approvals.list_query(w.event.id, status="pending", q=None, after=None, limit=50))
    nxt = await plan(approvals.list_query(w.event.id, status="pending", q=None, after=(NOW, ids[10]), limit=50))
    counts = await plan(select(R.status, func.count()).where(R.event_id == w.event.id).group_by(R.status))
    lock = await plan(
        select(R).where(R.id.in_(ids[:100]), R.event_id == w.event.id, R.status == "pending").order_by(R.submitted_at, R.id)
        .with_for_update(skip_locked=True)
    )
    decided = await plan(approvals.list_query(w.event.id, status="declined", q=None, after=None, limit=50))
    print(
        "\n--- pending page\n" + page + "\n--- next pending page\n" + nxt + "\n--- declined page\n" + decided
        + "\n--- counts\n" + counts + "\n--- approve lock\n" + lock
    )
    assert "ix_registration_requests_pending" in page and "Sort" not in page
    assert "ix_registration_requests_pending" in nxt and "Sort" not in nxt
    assert "ix_registration_requests_queue" in decided and "Sort" not in decided
    assert "ix_registration_requests_queue" in counts
    assert "Seq Scan" not in lock and ("ix_registration_requests_pending" in lock or "registration_requests_pkey" in lock)
