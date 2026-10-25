"""What the club dashboard's event screens lost on the way to the server.

- An edit that removed every interest (the form then leaves interest_ids
  out) failed instead of saving the event with none.
- The event's label (event_tag) was written on create only, and the event's
  info never returned it, so the edit form always showed it empty.
- The club's events list had no ticket price, so every paid event read "Free".
- The guest list's Paid / Unpaid filters, and the export's, were accepted and
  ignored: both gave every guest.
"""
import io
import json
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app.api.events.models import EventCategories
from app.api.interests.models import InterestCategory, Interests
from tests.conftest import auth_header


@pytest.fixture
async def world(session, make_club, mint_token):
    owner, club = await make_club(name="Edit Club", email="editclub@example.com")
    category = EventCategories(name="Edits")
    icat = InterestCategory(name="Edit interests")
    session.add_all([category, icat])
    await session.flush()
    interests = [Interests(name="Edit AI", category_id=icat.id), Interests(name="Edit Art", category_id=icat.id)]
    session.add_all(interests)
    await session.commit()
    form = {
        "name": "Edit Night",
        "event_datetime": (datetime.now(timezone.utc) + timedelta(days=9)).isoformat(), "duration": "2",
        "has_fee": "false", "status": "published", "category_id": str(category.id),
    }
    return dict(h=auth_header(await mint_token(owner)), form=form, interests=interests, club=club, owner=owner)


async def _create(client, world, **extra):
    made = await client.post("/api/v1/events/create", data={**world["form"], **extra}, headers=world["h"])
    assert made.status_code in (200, 201), made.text
    return made.json()["id"]


async def _info(client, event_id) -> dict:
    r = await client.get(f"/api/v1/events/info/{event_id}")
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.parametrize("sent", [None, "", "[]"], ids=["left-out", "empty", "empty-json"])
async def test_an_edit_that_removes_every_interest_saves_with_none(client, world, sent):
    ids = ",".join(str(i.id) for i in world["interests"])
    event_id = await _create(client, world, interest_ids=ids)
    assert len((await _info(client, event_id))["interests"]) == 2

    form = {**world["form"], "name": "Edit Night Again"}  # interest_ids as the case says
    if sent is not None:
        form["interest_ids"] = sent
    r = await client.put(f"/api/v1/events/update/{event_id}", data=form, headers=world["h"])
    assert r.status_code == 200, r.text
    info = await _info(client, event_id)
    assert info["name"] == "Edit Night Again"
    assert info["interests"] in ([], None)


async def test_interests_are_still_read_as_ids_or_a_json_array(client, world):
    a, b = world["interests"]
    event_id = await _create(client, world)
    for sent in (f"{a.id},{b.id}", json.dumps([a.id, b.id])):
        r = await client.put(f"/api/v1/events/update/{event_id}", data={**world["form"], "interest_ids": sent}, headers=world["h"])
        assert r.status_code == 200, r.text
        assert sorted(i["id"] for i in (await _info(client, event_id))["interests"]) == sorted([a.id, b.id])
    bad = await client.put(f"/api/v1/events/update/{event_id}", data={**world["form"], "interest_ids": "art,music"}, headers=world["h"])
    assert bad.status_code == 400


async def test_the_label_is_saved_on_edit_and_shown_to_the_form(client, world):
    event_id = await _create(client, world, event_tag="Workshop")
    assert (await _info(client, event_id))["event_tag"] == "Workshop"

    changed = await client.put(f"/api/v1/events/update/{event_id}", data={**world["form"], "event_tag": "Members only"}, headers=world["h"])
    assert changed.status_code == 200, changed.text
    assert (await _info(client, event_id))["event_tag"] == "Members only"

    # Emptied, the form leaves it out, and it is cleared like every other field.
    cleared = await client.put(f"/api/v1/events/update/{event_id}", data=world["form"], headers=world["h"])
    assert cleared.status_code == 200, cleared.text
    assert (await _info(client, event_id))["event_tag"] is None

    too_long = await client.put(f"/api/v1/events/update/{event_id}", data={**world["form"], "event_tag": "x" * 51}, headers=world["h"])
    assert too_long.status_code == 422


async def test_the_club_events_list_carries_the_ticket_price(client, world):
    paid = await _create(client, world, name="Paid Night", has_fee="true", reg_fee="250")
    free = await _create(client, world, name="Free Night")
    r = await client.get("/api/v1/clubs/events/list", headers=world["h"])
    assert r.status_code == 200, r.text
    by_id = {e["id"]: e for e in r.json()["items"]}
    assert by_id[paid]["has_fee"] is True and by_id[paid]["reg_fee"] == 250
    assert by_id[free]["has_fee"] is False and by_id[free]["reg_fee"] in (None, 0)


async def test_the_paid_and_unpaid_filters_apply_to_the_list_and_the_export(client, session, world, make_event, make_user, make_registration):
    event = await make_event(club=world["club"], created_by=world["owner"], name="Ticketed", has_fee=True, reg_fee=100)
    payer = await make_user(email="payer@example.com")
    owes = await make_user(email="owes@example.com")
    await make_registration(event=event, user=payer, is_paid=True, paid_amount=100, actual_amount=100)
    await make_registration(event=event, user=owes, is_paid=False, actual_amount=100)
    await session.commit()

    async def listed(**query):
        r = await client.get(f"/api/v1/events/registration/{event.id}/list", params=query, headers=world["h"])
        assert r.status_code == 200, r.text
        return sorted(i["email"] for i in r.json()["items"])

    async def exported(**query):
        r = await client.get(f"/api/v1/events/registration/{event.id}/export", params=query, headers=world["h"])
        assert r.status_code == 200, r.text
        sheet = load_workbook(io.BytesIO(r.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        email = rows[0].index("Email")
        return sorted(row[email] for row in rows[1:])

    everyone = ["owes@example.com", "payer@example.com"]
    assert await listed() == everyone
    assert await listed(is_paid="true") == ["payer@example.com"]
    assert await listed(is_paid="false") == ["owes@example.com"]
    assert await exported() == everyone
    assert await exported(is_paid="true") == ["payer@example.com"]
    assert await exported(is_paid="false") == ["owes@example.com"]
